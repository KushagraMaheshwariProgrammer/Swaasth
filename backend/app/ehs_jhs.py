"""Employees Health Scheme (EHS) and Journalists Health Scheme (JHS) shared engine."""

from __future__ import annotations

import csv
import json
import logging
import os
import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

logger = logging.getLogger(__name__)

_BACKEND_ROOT = Path(__file__).resolve().parent.parent
_DATA_DIR = _BACKEND_ROOT / "data" / "EHS_2017"
_PROCESSED_DIR = _DATA_DIR / "processed"

HOSPITAL_STRONG_MATCH = 0.82
HOSPITAL_CANDIDATE_MATCH = 0.55
PACKAGE_STRONG_MATCH = 0.72
PACKAGE_CANDIDATE_MATCH = 0.50

EHS_DISCLAIMER = (
    "Employees Health Scheme is intended for eligible Telangana government "
    "employees, pensioners, and dependent family members. Final eligibility, "
    "card validity, package approval, and hospital empanelment must be verified "
    "through official EHF/Aarogyasri sources."
)
JHS_DISCLAIMER = (
    "Journalists Health Scheme is intended for eligible working/retired "
    "journalists and dependent family members. EHS network hospitals serve JHS, "
    "but final eligibility, card validity, package approval, and hospital "
    "empanelment must be verified through official EHF/JHS/Aarogyasri sources."
)

_PACKAGE_NOT_MATCHED = (
    "EHS/JHS package was not matched from the available dataset. "
    "CGHS fallback is shown only as a general benchmark."
)
_HOSPITAL_FOUND = "Hospital found in EHS/JHS empanelled hospital data."
_HOSPITAL_NOT_FOUND = (
    "Hospital was not found in EHS/JHS empanelled hospital data. "
    "Manual verification with EHF/Aarogyasri sources is required."
)
_DIRECT_PAYMENT = (
    "EHS/JHS are cashless package-based schemes for eligible covered treatments. "
    "Direct collection from the patient may require manual verification."
)

_DIRECT_PAYMENT_KEYWORDS = (
    "amount paid",
    "cash paid",
    "paid by patient",
    "self pay",
    "direct payment",
    "payment received",
    "advance paid",
    "balance paid",
    "cash payment",
)

_STRIP_TOKENS = frozenset(
    {
        "hospital",
        "hospitals",
        "hosp",
        "medical",
        "centre",
        "center",
        "institute",
        "pvt",
        "ltd",
        "llp",
        "trust",
        "and",
        "the",
        "of",
        "a",
        "private",
        "limited",
        "care",
        "nursing",
        "home",
    }
)


def _dev_logging_enabled() -> bool:
    return os.getenv("ENV", "").lower() in {"dev", "development", "local"} or os.getenv(
        "EHS_JHS_DEBUG", ""
    ).lower() in {"1", "true", "yes"}


def _clean_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip())


def normalize_text(text: str) -> str:
    normalized = (text or "").lower().strip()
    normalized = normalized.replace("&", " and ")
    normalized = re.sub(r"[^a-z0-9 ]+", " ", normalized)
    normalized = _clean_spaces(normalized)
    tokens = [token for token in normalized.split() if token not in _STRIP_TOKENS]
    return _clean_spaces(" ".join(tokens))


def build_aliases(name: str, extra_aliases: Iterable[str] | None = None) -> set[str]:
    aliases: set[str] = set()
    raw = _clean_spaces(name)
    if not raw:
        return aliases
    aliases.add(normalize_text(raw))
    words = [word for word in raw.split() if word.lower() not in _STRIP_TOKENS]
    if words:
        aliases.add(normalize_text(" ".join(words)))
    if extra_aliases:
        for alias in extra_aliases:
            cleaned = _clean_spaces(str(alias))
            if cleaned:
                aliases.add(normalize_text(cleaned))
    aliases.discard("")
    return aliases


def _parse_amount(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).replace(",", "").strip()
    if not text:
        return None
    try:
        amount = float(text)
    except ValueError:
        return None
    return amount if amount >= 0 else None


def parse_ehs_speciality_names(raw: str) -> list[str]:
    if not raw or not raw.strip():
        return []
    names: list[str] = []
    seen: set[str] = set()
    for match in re.finditer(
        r"\b([SM]\d+)\s*-\s*(.+?)(?=\s+[SM]\d+\s*-|\s*$)",
        raw,
        flags=re.IGNORECASE,
    ):
        label = _clean_spaces(match.group(2))
        if not label:
            continue
        key = label.casefold()
        if key in seen:
            continue
        seen.add(key)
        names.append(label.title() if label == label.upper() else label)
    return names


@dataclass
class EhsJhsHospitalRecord:
    hospital_name: str
    normalized_name: str
    hospital_type: str
    district: str
    mandal: str
    municipality: str
    state: str
    empanelment_status: str
    specialities_raw: str
    specialities_readable: tuple[str, ...]
    source: str
    alias_keys: set[str] = field(default_factory=set)


@dataclass
class EhsJhsPackageRecord:
    package_code: str
    package_name: str
    normalized_name: str
    specialty_code: str
    category: str
    approved_rate: float
    rate_private_nabh: float | None
    rate_private_non_nabh: float | None
    rate_semi_private_nabh: float | None
    rate_semi_private_non_nabh: float | None
    source_file: str
    source_sheet: str
    tokens: set[str] = field(default_factory=set)


class EhsJhsStore:
    def __init__(self) -> None:
        self.hospitals: list[EhsJhsHospitalRecord] = []
        self.packages: list[EhsJhsPackageRecord] = []
        self.meta: dict[str, Any] = {}
        self._hospital_alias_index: dict[str, list[EhsJhsHospitalRecord]] = {}
        self._hospital_name_index: dict[str, list[EhsJhsHospitalRecord]] = {}
        self._package_name_index: dict[str, list[EhsJhsPackageRecord]] = {}
        self._load()

    def _load(self) -> None:
        loaded_packages = self._load_packages()
        loaded_hospitals = self._load_hospitals()
        if not loaded_packages:
            logger.warning("EHS/JHS package data unavailable")
        if not loaded_hospitals:
            logger.warning("EHS/JHS hospital data unavailable")
        if _dev_logging_enabled():
            logger.info(
                "EHS/JHS loaded packages=%d hospitals=%d primary=%s sheet=%s",
                len(self.packages),
                len(self.hospitals),
                self.meta.get("primary_file"),
                self.meta.get("primary_sheet"),
            )

    def _load_packages(self) -> bool:
        processed = _PROCESSED_DIR / "packages.json"
        if processed.exists():
            payload = json.loads(processed.read_text(encoding="utf-8"))
            self.meta = payload.get("meta") or {}
            for raw in payload.get("packages") or []:
                self._append_package(raw)
            return bool(self.packages)

        primary = _DATA_DIR / "EHS_with_Category_Hierarchy.xlsx"
        fallback = _DATA_DIR / "EHS-20-12-14_high_quality.xlsx"
        path = primary if primary.exists() else fallback
        if not path.exists():
            return False
        try:
            import openpyxl
        except ImportError:
            logger.warning("openpyxl not installed; EHS/JHS packages unavailable")
            return False

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet_name = workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return False
        headers = [str(h or "").strip() for h in rows[0]]
        hmap = {h.lower(): i for i, h in enumerate(headers)}

        def idx(name: str) -> int | None:
            return hmap.get(name.lower())

        self.meta = {
            "primary_file": path.name,
            "primary_sheet": sheet_name,
            "columns_mapped": {
                "package_code": "Procedure_ICD_Code",
                "package_name": "Procedure_Name",
                "specialty": "Applicable_Speciality_Code",
                "category": "Category_Hierarchy",
                "approved_rate": "Price_Private_NABH",
            },
        }

        seen: set[str] = set()
        for row in rows[1:]:
            name_i = idx("procedure_name")
            if name_i is None or not row or len(row) <= name_i:
                continue
            name = str(row[name_i] or "").strip()
            if not name:
                continue
            code_i = idx("procedure_icd_code")
            code = str(row[code_i] or "").strip() if code_i is not None else ""
            dedupe = (code or name).casefold()
            if dedupe in seen:
                continue
            seen.add(dedupe)
            spec_i = idx("applicable_speciality_code")
            cat_i = idx("category_hierarchy")
            rate_nabh = _parse_amount(
                row[idx("price_private_nabh")] if idx("price_private_nabh") is not None else None
            )
            rate_non_nabh = _parse_amount(
                row[idx("price_private_non_nabh")]
                if idx("price_private_non_nabh") is not None
                else None
            )
            rate_semi_nabh = _parse_amount(
                row[idx("price_semi_private_nabh")]
                if idx("price_semi_private_nabh") is not None
                else None
            )
            rate_semi_non = _parse_amount(
                row[idx("price_semi_private_non_nabh")]
                if idx("price_semi_private_non_nabh") is not None
                else None
            )
            approved = rate_nabh or rate_non_nabh or rate_semi_nabh or rate_semi_non
            if approved is None:
                continue
            self._append_package(
                {
                    "package_code": code or name,
                    "package_name": name,
                    "specialty_code": str(row[spec_i] or "").strip() if spec_i is not None else "",
                    "category": str(row[cat_i] or "").strip() if cat_i is not None else "",
                    "approved_rate": approved,
                    "rate_private_nabh": rate_nabh,
                    "rate_private_non_nabh": rate_non_nabh,
                    "rate_semi_private_nabh": rate_semi_nabh,
                    "rate_semi_private_non_nabh": rate_semi_non,
                    "source_file": path.name,
                    "source_sheet": sheet_name,
                }
            )
        return bool(self.packages)

    def _append_package(self, raw: dict[str, Any]) -> None:
        name = str(raw.get("package_name") or "").strip()
        if not name:
            return
        normalized = normalize_text(name)
        record = EhsJhsPackageRecord(
            package_code=str(raw.get("package_code") or name).strip(),
            package_name=name,
            normalized_name=normalized,
            specialty_code=str(raw.get("specialty_code") or "").strip(),
            category=str(raw.get("category") or "").strip(),
            approved_rate=float(raw.get("approved_rate") or 0),
            rate_private_nabh=_parse_amount(raw.get("rate_private_nabh")),
            rate_private_non_nabh=_parse_amount(raw.get("rate_private_non_nabh")),
            rate_semi_private_nabh=_parse_amount(raw.get("rate_semi_private_nabh")),
            rate_semi_private_non_nabh=_parse_amount(raw.get("rate_semi_private_non_nabh")),
            source_file=str(raw.get("source_file") or ""),
            source_sheet=str(raw.get("source_sheet") or ""),
            tokens=set(normalized.split()),
        )
        self.packages.append(record)
        self._package_name_index.setdefault(normalized, []).append(record)

    def _load_hospitals(self) -> bool:
        processed = _PROCESSED_DIR / "hospitals.json"
        if processed.exists():
            payload = json.loads(processed.read_text(encoding="utf-8"))
            for raw in payload.get("hospitals") or []:
                self._append_hospital(raw)
            return bool(self.hospitals)

        csv_path = _DATA_DIR / "EHS_hospital_list.csv"
        if not csv_path.exists():
            return False
        seen: set[str] = set()
        with csv_path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            for raw in reader:
                name = (raw.get("Hospital Name") or raw.get("hospital_name") or "").strip()
                if not name:
                    continue
                key = name.casefold()
                if key in seen:
                    continue
                seen.add(key)
                self._append_hospital(
                    {
                        "hospital_name": name,
                        "hospital_type": (raw.get("Hospital Type") or "").strip(),
                        "district": (raw.get("District") or "").strip(),
                        "mandal": (raw.get("Mandal") or "").strip(),
                        "municipality": (raw.get("Municipality") or "").strip(),
                        "specialities_raw": (raw.get("Specialties") or "").strip(),
                        "empanelment_status": "active",
                        "source": "EHS_hospital_list.csv",
                    }
                )
        return bool(self.hospitals)

    def _append_hospital(self, raw: dict[str, Any]) -> None:
        name = str(raw.get("hospital_name") or "").strip()
        if not name:
            return
        normalized = normalize_text(name)
        readable = parse_ehs_speciality_names(str(raw.get("specialities_raw") or ""))
        record = EhsJhsHospitalRecord(
            hospital_name=name,
            normalized_name=normalized,
            hospital_type=str(raw.get("hospital_type") or "").strip(),
            district=str(raw.get("district") or "").strip(),
            mandal=str(raw.get("mandal") or "").strip(),
            municipality=str(raw.get("municipality") or "").strip(),
            state="Telangana",
            empanelment_status=str(raw.get("empanelment_status") or "active").strip(),
            specialities_raw=str(raw.get("specialities_raw") or "").strip(),
            specialities_readable=tuple(readable),
            source=str(raw.get("source") or "EHS_hospital_list.csv"),
            alias_keys=build_aliases(name),
        )
        self.hospitals.append(record)
        self._hospital_name_index.setdefault(normalized, []).append(record)
        for alias in record.alias_keys:
            self._hospital_alias_index.setdefault(alias, []).append(record)

    def _location_bonus(
        self,
        record: EhsJhsHospitalRecord,
        *,
        patient_district: str | None,
        patient_city: str | None,
    ) -> float:
        bonus = 0.0
        if patient_district and record.district:
            if normalize_text(patient_district) == normalize_text(record.district):
                bonus += 0.08
        if patient_city:
            city_norm = normalize_text(patient_city)
            if record.municipality and city_norm == normalize_text(record.municipality):
                bonus += 0.05
            elif record.mandal and city_norm == normalize_text(record.mandal):
                bonus += 0.04
        return bonus

    def match_hospital(
        self,
        ocr_hospital_name: str | None,
        *,
        patient_district: str | None = None,
        patient_city: str | None = None,
    ) -> dict[str, Any]:
        ocr_name = _clean_spaces(ocr_hospital_name or "")
        base = {
            "ocr_hospital_name": ocr_name,
            "matched_hospital_name": None,
            "confidence_score": 0.0,
            "match_reason": "No hospital match found in EHS/JHS directory",
            "district": patient_district,
            "mandal": None,
            "municipality": patient_city,
            "hospital_type": None,
            "empanelment_status": None,
            "specialities": None,
            "specialities_readable": [],
            "source_file": "EHS_hospital_list.csv",
            "source_type": "ehs_jhs_primary",
            "status": "not_found",
            "candidates": [],
        }
        if not ocr_name:
            base["match_reason"] = "OCR hospital name missing"
            base["status"] = "manual_verification_required"
            return base
        if not self.hospitals:
            base["match_reason"] = "EHS/JHS hospital database unavailable"
            base["status"] = "manual_verification_required"
            return base

        query_norm = normalize_text(ocr_name)
        query_aliases = build_aliases(ocr_name)
        candidates: dict[int, tuple[float, str, EhsJhsHospitalRecord]] = {}

        for alias in query_aliases:
            for record in self._hospital_alias_index.get(alias, []):
                score = 1.0 + self._location_bonus(
                    record,
                    patient_district=patient_district,
                    patient_city=patient_city,
                )
                candidates[id(record)] = (min(score, 1.0), "Exact alias match", record)

        for record in self._hospital_name_index.get(query_norm, []):
            score = 1.0 + self._location_bonus(
                record,
                patient_district=patient_district,
                patient_city=patient_city,
            )
            candidates[id(record)] = (min(score, 1.0), "Exact normalized name match", record)

        query_tokens = set(query_norm.split())
        for record in self.hospitals:
            if id(record) in candidates:
                continue
            ratio = SequenceMatcher(None, query_norm, record.normalized_name).ratio()
            if query_tokens and record.normalized_name:
                overlap = len(query_tokens & set(record.normalized_name.split())) / max(
                    len(query_tokens), 1
                )
                ratio = max(ratio, overlap)
            if query_norm in record.normalized_name or record.normalized_name in query_norm:
                ratio = max(ratio, 0.9)
            ratio += self._location_bonus(
                record,
                patient_district=patient_district,
                patient_city=patient_city,
            )
            if ratio >= HOSPITAL_CANDIDATE_MATCH:
                reason = (
                    "Contains match"
                    if query_norm in record.normalized_name
                    or record.normalized_name in query_norm
                    else "Fuzzy name match"
                )
                candidates[id(record)] = (min(ratio, 1.0), reason, record)

        if not candidates:
            return base

        scored = sorted(candidates.values(), key=lambda item: item[0], reverse=True)
        best_score, best_reason, best_record = scored[0]
        if best_score < HOSPITAL_STRONG_MATCH and len(scored) > 1:
            return {
                **base,
                "confidence_score": round(best_score, 3),
                "match_reason": "Multiple possible EHS/JHS hospital matches",
                "status": "manual_verification_required",
                "candidates": [
                    {
                        "matched_hospital_name": rec.hospital_name,
                        "confidence_score": round(score, 3),
                        "district": rec.district,
                        "municipality": rec.municipality,
                    }
                    for score, _, rec in scored[:6]
                ],
            }

        status = "empanelled" if best_score >= HOSPITAL_STRONG_MATCH else "manual_verification_required"
        readable = list(best_record.specialities_readable)
        return {
            "ocr_hospital_name": ocr_name,
            "matched_hospital_name": best_record.hospital_name,
            "confidence_score": round(best_score, 3),
            "match_reason": best_reason,
            "district": best_record.district or patient_district,
            "mandal": best_record.mandal or None,
            "municipality": best_record.municipality or patient_city,
            "hospital_type": best_record.hospital_type or None,
            "empanelment_status": best_record.empanelment_status,
            "specialities": ", ".join(readable) if readable else None,
            "specialities_readable": readable,
            "source_file": "EHS_hospital_list.csv",
            "source_type": "ehs_jhs_primary",
            "status": status,
            "candidates": [],
        }

    def match_package(self, ocr_item_name: str | None) -> dict[str, Any]:
        query = _clean_spaces(ocr_item_name or "")
        if not query or not self.packages:
            return {
                "matched_package": None,
                "package_code": None,
                "package_name": None,
                "specialty": None,
                "category_or_specialty": None,
                "approved_rate": None,
                "source_file": self.meta.get("primary_file"),
                "source_sheet": self.meta.get("primary_sheet"),
                "confidence_score": 0.0,
                "match_reason": "Package not matched",
            }

        query_norm = normalize_text(query)
        best: tuple[float, str, EhsJhsPackageRecord] | None = None
        for record in self.packages:
            if query_norm == record.normalized_name:
                best = (1.0, "Exact package name match", record)
                break
            ratio = SequenceMatcher(None, query_norm, record.normalized_name).ratio()
            query_tokens = set(query_norm.split())
            if query_tokens and record.tokens:
                overlap = len(query_tokens & record.tokens) / max(len(query_tokens), 1)
                ratio = max(ratio, overlap)
            if query_norm in record.normalized_name or record.normalized_name in query_norm:
                ratio = max(ratio, 0.88)
            if ratio >= PACKAGE_CANDIDATE_MATCH and (
                best is None or ratio > best[0]
            ):
                reason = (
                    "Contains match"
                    if query_norm in record.normalized_name
                    or record.normalized_name in query_norm
                    else "Fuzzy package match"
                )
                best = (ratio, reason, record)

        if best is None:
            return {
                "matched_package": None,
                "package_code": None,
                "package_name": None,
                "specialty": None,
                "category_or_specialty": None,
                "approved_rate": None,
                "source_file": self.meta.get("primary_file"),
                "source_sheet": self.meta.get("primary_sheet"),
                "confidence_score": 0.0,
                "match_reason": "Package not matched",
            }

        score, reason, record = best
        category_or_specialty = record.category or record.specialty_code or None
        return {
            "matched_package": record.package_name,
            "package_code": record.package_code,
            "package_name": record.package_name,
            "specialty": record.specialty_code,
            "category_or_specialty": category_or_specialty,
            "approved_rate": record.approved_rate,
            "source_file": record.source_file,
            "source_sheet": record.source_sheet,
            "confidence_score": round(score, 3),
            "match_reason": reason,
        }


_store: EhsJhsStore | None = None


def get_ehs_jhs_store() -> EhsJhsStore:
    global _store
    if _store is None:
        _store = EhsJhsStore()
    return _store


def _bill_shows_direct_payment(*, ocr_text: str | None, line_items: list[dict[str, Any]] | None) -> bool:
    haystack = (ocr_text or "").lower()
    if any(keyword in haystack for keyword in _DIRECT_PAYMENT_KEYWORDS):
        return True
    for item in line_items or []:
        name = str(item.get("item_name") or "").lower()
        if any(keyword in name for keyword in ("advance", "deposit", "cash paid", "self pay")):
            return True
    return False


def _comparison_status(charged: float, approved: float | None, *, matched: bool) -> str:
    if not matched or approved is None:
        return "Package Not Found"
    if charged > approved:
        return "Excess Over Package Rate"
    return "Within Package Rate"


def build_ehs_jhs_report(
    *,
    scheme_type: str,
    selected: bool,
    ehs_is_government_employee: bool = False,
    ehs_is_pensioner: bool = False,
    ehs_is_dependent: bool = False,
    ehs_has_health_card: bool = False,
    ehs_card_number: str | None = None,
    jhs_is_working_journalist: bool = False,
    jhs_is_retired_journalist: bool = False,
    jhs_is_dependent: bool = False,
    jhs_has_health_card: bool = False,
    jhs_has_aadhaar: bool = False,
    jhs_card_number: str | None = None,
    ocr_hospital_name: str | None = None,
    patient_district: str | None = None,
    patient_city: str | None = None,
    line_items: list[dict[str, Any]] | None = None,
    compared_line_items: list[dict[str, Any]] | None = None,
    ocr_text: str | None = None,
) -> dict[str, Any] | None:
    if not selected:
        return None

    scheme = "EHS" if scheme_type.upper() == "EHS" else "JHS"
    try:
        store = get_ehs_jhs_store()
    except Exception as exc:
        logger.exception("EHS/JHS store unavailable: %s", exc)
        return {
            "scheme_type": scheme,
            "selected": True,
            "status": "manual_verification_required",
            "eligibility_snapshot": {},
            "hospital_verification": {
                "status": "manual_verification_required",
                "match_reason": str(exc),
            },
            "package_comparisons": [],
            "advisories": ["EHS/JHS data is unavailable. Manual verification required."],
            "disclaimer": EHS_DISCLAIMER if scheme == "EHS" else JHS_DISCLAIMER,
        }

    hospital_verification = store.match_hospital(
        ocr_hospital_name,
        patient_district=patient_district,
        patient_city=patient_city,
    )

    package_comparisons: list[dict[str, Any]] = []
    items = line_items or []
    compared = compared_line_items or []
    package_not_matched = False

    for index, item in enumerate(items):
        bill_item_name = str(item.get("item_name", "")).strip()
        charged_amount = _parse_amount(item.get("total_price")) or 0.0
        package_match = store.match_package(bill_item_name)
        matched = bool(package_match.get("matched_package"))
        fallback_used = not matched
        approved_rate = package_match.get("approved_rate")
        cghs_item = compared[index] if index < len(compared) else {}
        if fallback_used:
            package_not_matched = True
            approved_rate = _parse_amount(cghs_item.get("cghs_rate"))
        excess_amount = None
        if approved_rate is not None:
            excess_amount = round(max(charged_amount - approved_rate, 0.0), 2)
        if matched and (package_match.get("confidence_score") or 0) < PACKAGE_STRONG_MATCH:
            status = "Manual Verification Required"
        elif fallback_used:
            status = "CGHS Fallback Used" if approved_rate is not None else "Package Not Found"
        else:
            status = _comparison_status(charged_amount, approved_rate, matched=True)

        package_comparisons.append(
            {
                "bill_item_name": bill_item_name,
                "charged_amount": charged_amount,
                "matched_package_code": package_match.get("package_code"),
                "matched_package_name": package_match.get("package_name"),
                "category_or_specialty": package_match.get("category_or_specialty"),
                "approved_rate": approved_rate,
                "source_file": package_match.get("source_file"),
                "source_sheet": package_match.get("source_sheet"),
                "confidence_score": package_match.get("confidence_score"),
                "match_reason": package_match.get("match_reason"),
                "excess_amount": excess_amount,
                "status": status,
                "fallback_used": fallback_used,
            }
        )

    advisories: list[str] = []
    if scheme == "EHS":
        advisories.append(EHS_DISCLAIMER)
    else:
        advisories.append(JHS_DISCLAIMER)

    hospital_status = str(hospital_verification.get("status") or "")
    if hospital_status == "empanelled":
        advisories.append(_HOSPITAL_FOUND)
    else:
        advisories.append(_HOSPITAL_NOT_FOUND)

    if package_not_matched:
        advisories.append(_PACKAGE_NOT_MATCHED)

    if not (ehs_has_health_card if scheme == "EHS" else jhs_has_health_card):
        advisories.append(
            "Patient health card status was not confirmed. Verify EHS/JHS card validity separately."
        )

    if hospital_status == "empanelled" and _bill_shows_direct_payment(
        ocr_text=ocr_text, line_items=items
    ):
        advisories.append(_DIRECT_PAYMENT)

    report_status = "matched"
    if (
        hospital_status != "empanelled"
        or package_not_matched
        or any(
            row.get("status") in {"Manual Verification Required", "Package Not Found", "CGHS Fallback Used"}
            for row in package_comparisons
        )
    ):
        report_status = "manual_verification_required"

    eligibility_snapshot = (
        {
            "government_employee": ehs_is_government_employee,
            "pensioner": ehs_is_pensioner,
            "dependent": ehs_is_dependent,
            "has_health_card": ehs_has_health_card,
            "card_number": ehs_card_number,
        }
        if scheme == "EHS"
        else {
            "working_journalist": jhs_is_working_journalist,
            "retired_journalist": jhs_is_retired_journalist,
            "dependent": jhs_is_dependent,
            "has_health_card": jhs_has_health_card,
            "has_aadhaar": jhs_has_aadhaar,
            "card_number": jhs_card_number,
        }
    )

    return {
        "scheme_type": scheme,
        "selected": True,
        "status": report_status,
        "eligibility_snapshot": eligibility_snapshot,
        "hospital_verification": hospital_verification,
        "package_comparisons": package_comparisons,
        "advisories": advisories,
        "disclaimer": EHS_DISCLAIMER if scheme == "EHS" else JHS_DISCLAIMER,
    }
